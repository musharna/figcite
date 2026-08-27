"""`figcite bib` dispatches on a suffix, twice, and only one side was probed.

    _records_for_bib  `if p.suffix.lower() == ".json":`   parse as a manifest
    cmd_bib           `if ... p.suffix.lower() != ".json":` derive <source>.bib

Both compare against ".json", and every source any test has ever passed is a
`.pptx`, a `.pdf` or a `.json` -- and `.pdf` and `.pptx` both sort ABOVE
".json". So nothing has ever been below it, and both narrowings survive:

    `== ".json"` -> `<=`   claims everything below, and json.loads a deck
    `!= ".json"` -> `>`    stops deriving an output path for anything below,
                           so the .bib is printed to stdout and never written

A suffix below ".json" is not exotic -- ".csv" is the other half of the very
manifest this command reads, and ".bib" is its own output. This uses ".dat"
because it has to be a suffix the dispatcher will actually HANDLE: the
extension routes it away from the JSON and PDF arms, and python-pptx opens a
deck by sniffing the zip rather than by trusting the name, so one real deck
under an odd extension exercises both comparisons at once.
"""

from __future__ import annotations

import pytest
from PIL import Image

from pptx import Presentation
from pptx.util import Inches

from figcite import cli, store
from figcite.provenance import Record, embed

CITE = "Shiragaki et al. 2020"


@pytest.fixture
def deck_named_dat(tmp_path):
    """A real .pptx whose filename sorts below '.json'."""
    raw = tmp_path / "raw.png"
    Image.new("RGB", (60, 45), (20, 80, 120)).save(raw)
    img = tmp_path / "fig.png"
    rec = embed(
        raw,
        img,
        Record(citation=CITE, short_cite=CITE, confirmed=True, doi="10.1/x"),
    )
    store.put(rec)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(img), Inches(0.5), Inches(0.5), width=Inches(3))
    path = tmp_path / "deck.dat"
    prs.save(str(path))
    return path


def test_the_fixture_suffix_really_sorts_below_json():
    """The premise. `.pptx` and `.pdf` -- everything the existing tests pass --
    sort ABOVE '.json', which is why neither narrowing was observable."""
    assert ".dat" < ".json", ".dat"
    assert ".csv" < ".json", ".csv"
    assert ".bib" < ".json", ".bib"
    assert ".pdf" > ".json" and ".pptx" > ".json", "the usual sources"


def test_a_deck_under_an_odd_suffix_is_not_parsed_as_a_manifest(deck_named_dat):
    """`== ".json"` widened to `<=` hands a zip file to json.loads."""
    from figcite.cli import _records_for_bib

    recs = _records_for_bib(str(deck_named_dat))

    assert len(recs) == 1, recs
    assert recs[0].doi == "10.1/x", recs[0]


def test_a_json_manifest_is_still_parsed_as_one(tmp_path):
    """Positive control. "Never treat anything as JSON" satisfies the test
    above and removes the manifest arm entirely."""
    import json

    from figcite.cli import _records_for_bib

    manifest = tmp_path / "man.json"
    manifest.write_text(
        json.dumps([{"record": {"doi": "10.1/from-manifest", "citation": CITE}}]),
        encoding="utf-8",
    )

    recs = _records_for_bib(str(manifest))

    assert [r.doi for r in recs] == ["10.1/from-manifest"], recs


def test_a_bib_is_written_next_to_a_source_with_an_odd_suffix(deck_named_dat, capsys):
    """`!= ".json"` widened to `>` stops deriving an output path for anything
    sorting below, so the BibTeX goes to stdout and no file is written -- and
    `figcite bib` exists to produce a file for ghostcite to check."""
    rc = cli.main(["bib", str(deck_named_dat)])
    assert rc == 0, capsys.readouterr()

    expected = deck_named_dat.with_suffix("").with_suffix(".bib")
    assert expected.exists(), (
        f"no .bib was written beside {deck_named_dat.name}; the output path is "
        f"only derived for suffixes that sort one way"
    )
    assert "10.1/x" in expected.read_text(encoding="utf-8")


def test_a_json_source_still_prints_rather_than_writing(tmp_path, capsys):
    """The other half of that guard, and the reason it exists: a manifest's
    own name would make `<manifest>.bib` collide confusingly, so the JSON arm
    prints instead."""
    import json

    manifest = tmp_path / "man2.json"
    manifest.write_text(
        # `confirmed` matters: records_to_bibtex drops unconfirmed
        # figure-to-DOI links by default, which is the ghost-citation guard
        # and not what this test is about.
        json.dumps(
            [
                {
                    "record": {
                        "doi": "10.1/printed",
                        "citation": CITE,
                        "confirmed": True,
                    }
                }
            ]
        ),
        encoding="utf-8",
    )

    rc = cli.main(["bib", str(manifest)])
    out = capsys.readouterr().out

    assert rc == 0
    assert not (tmp_path / "man2.bib").exists(), "a .bib was written for a manifest"
    assert "10.1/printed" in out, out
