"""Deck writing, verified by reopening the .pptx and by unzipping it.

The unconfirmed-citation gate is tested in both directions in one test: the
guess must NOT reach the slide by default, and MUST reach it under the explicit
override. A one-directional test there would pass on a deck writer that simply
printed nothing at all.
"""

import zipfile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from figcite import store
from figcite.deck import apply, audit, get_alt_text, iter_pictures
from figcite.provenance import Record, embed, read_embedded


CITE = (
    "Shiragaki et al. (2020). Phylogenetic Analysis of Capsicum. Horticulturae 6: 87."
)
DOI = "10.3390/horticulturae6040087"


def _make_image(path, color, size=(400, 300)):
    im = Image.new("RGB", size)
    for x in range(size[0]):  # non-flat so dhash is meaningful
        for y in range(0, size[1], 3):
            im.putpixel((x, y), ((x + color) % 256, (y * 2) % 256, color % 256))
    im.save(path)
    return path


def _tagged(tmp_path, name, *, confirmed=True, doi=DOI, cite=CITE, color=7):
    src = _make_image(tmp_path / f"raw-{name}", color)
    rec = Record(
        doi=doi,
        citation=cite,
        short_cite="Shiragaki et al. 2020",
        license_url="https://creativecommons.org/licenses/by/4.0",
        reuse="reuse-ok-attribution-required",
        source_kind="pdf-crop",
        confirmed=confirmed,
    )
    out = tmp_path / name
    rec = embed(src, out, rec)
    store.put(rec)
    return out, rec


def _deck_with(tmp_path, images, name="deck.pptx"):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    top = Inches(0.5)
    for img in images:
        slide.shapes.add_picture(str(img), Inches(0.5), top, width=Inches(3))
        top += Inches(2.4)
    p = tmp_path / name
    prs.save(str(p))
    return p


def _all_text(pptx_path):
    prs = Presentation(str(pptx_path))
    out = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
    return "\n".join(out)


def test_alt_text_absent_before_and_present_after(tmp_path):
    img, rec = _tagged(tmp_path, "fig1.png")
    deck = _deck_with(tmp_path, [img])

    prs = Presentation(str(deck))
    pics = [p for s in prs.slides for p, _ in iter_pictures(s)]
    assert len(pics) == 1
    # Inserting a picture pre-fills alt text with the FILENAME (python-pptx and
    # PowerPoint both do this), so "no alt text" is not the pre-state -- "no
    # provenance in the alt text" is.
    before = get_alt_text(pics[0])
    assert before == "fig1.png", f"unexpected pre-state alt text: {before!r}"
    assert DOI not in before, "the DOI was already there; test proves nothing"

    out = tmp_path / "out.pptx"
    apply(deck, out, captions=True, credits=True, manifest_path=tmp_path / "man")

    prs2 = Presentation(str(out))
    pics2 = [p for s in prs2.slides for p, _ in iter_pictures(s)]
    alt = get_alt_text(pics2[0])
    assert DOI in alt, f"alt text lost the DOI: {alt!r}"
    assert "creativecommons" in alt, "alt text dropped the license"


def test_caption_and_credits_land_on_slides(tmp_path):
    img, _ = _tagged(tmp_path, "fig2.png")
    deck = _deck_with(tmp_path, [img])
    out = tmp_path / "out.pptx"
    rep = apply(deck, out, captions=True, credits=True, manifest_path=tmp_path / "man")

    text = _all_text(out)
    assert "[1]" in text, "no caption number reached the slide"
    assert "Shiragaki" in text, "no citation text reached the deck"
    assert "Image credits" in text, "credits slide missing"
    assert rep["cited"] == 1

    prs = Presentation(str(out))
    names = [sh.name for s in prs.slides for sh in s.shapes]
    assert any(n.startswith("figcite-caption") for n in names)
    assert any(n.startswith("figcite-credits-marker") for n in names)
    assert len(prs.slides.__iter__.__self__._sldIdLst) == 2  # original + credits


def test_media_bytes_keep_embedded_metadata_through_pptx(tmp_path):
    img, rec = _tagged(tmp_path, "fig3.png")
    deck = _deck_with(tmp_path, [img])
    out = tmp_path / "out.pptx"
    apply(deck, out, captions=False, credits=False, manifest_path=None)

    z = zipfile.ZipFile(str(out))
    media = [n for n in z.namelist() if n.startswith("ppt/media/")]
    assert media, "no media in the pptx"
    found = [read_embedded(z.read(n)) for n in media]
    found = [f for f in found if f is not None]
    assert found and found[0].doi == DOI, (
        "the DOI did not survive inside ppt/media -- layer 1 is dead"
    )


def test_unconfirmed_guess_is_withheld_then_released(tmp_path):
    """Both directions, because a writer that prints nothing would pass one."""
    img, _ = _tagged(tmp_path, "fig4.png", confirmed=False, color=99)
    deck = _deck_with(tmp_path, [img])

    strict = tmp_path / "strict.pptx"
    apply(deck, strict, captions=True, credits=True, allow_unconfirmed=False)
    t1 = _all_text(strict)
    assert "Shiragaki" not in t1, "an unconfirmed guess was printed as a citation"
    assert "unconfirmed" in t1.lower(), "the gap was hidden instead of flagged"

    loose = tmp_path / "loose.pptx"
    apply(deck, loose, captions=True, credits=True, allow_unconfirmed=True)
    t2 = _all_text(loose)
    assert "Shiragaki" in t2, (
        "--allow-unconfirmed printed nothing either; the gate is not a gate, it is a wall"
    )


def test_untagged_image_is_reported_not_hidden(tmp_path):
    tagged, _ = _tagged(tmp_path, "fig5.png", color=11)
    stranger = _make_image(tmp_path / "stranger.png", 200)
    deck = _deck_with(tmp_path, [tagged, stranger])

    rep_audit = audit(deck)
    assert rep_audit["untagged_substantive"] == 1

    out = tmp_path / "out.pptx"
    rep = apply(deck, out, captions=True, credits=True, manifest_path=tmp_path / "man")
    assert rep["unsourced"] == 1
    text = _all_text(out)
    assert "no recorded source" in text, (
        "the deck silently omitted an unsourced image instead of saying so"
    )

    prs = Presentation(str(out))
    pics = [p for s in prs.slides for p, _ in iter_pictures(s)]
    alts = [get_alt_text(p) for p in pics]
    assert any("not recorded" in a for a in alts)


def test_apply_is_idempotent(tmp_path):
    img, _ = _tagged(tmp_path, "fig6.png", color=31)
    deck = _deck_with(tmp_path, [img])
    once = tmp_path / "once.pptx"
    apply(deck, once, captions=True, credits=True)
    twice = tmp_path / "twice.pptx"
    rep = apply(once, twice, captions=True, credits=True)

    assert rep["removed_prior_figcite_shapes"] >= 2, (
        "prior figcite shapes were not cleaned up"
    )
    prs = Presentation(str(twice))
    caps = [
        sh.name
        for s in prs.slides
        for sh in s.shapes
        if sh.name.startswith("figcite-caption")
    ]
    credits = [
        sh.name
        for s in prs.slides
        for sh in s.shapes
        if sh.name.startswith("figcite-credits-marker")
    ]
    assert len(caps) == 1, f"captions duplicated on re-run: {caps}"
    assert len([c for c in credits if "body" not in c]) == 1, (
        f"credits slides duplicated: {credits}"
    )
    assert len(prs.slides.__iter__.__self__._sldIdLst) == 2


def test_manifest_csv_has_the_doi(tmp_path):
    img, _ = _tagged(tmp_path, "fig7.png", color=57)
    deck = _deck_with(tmp_path, [img])
    out = tmp_path / "out.pptx"
    rep = apply(deck, out, manifest_path=tmp_path / "manifest")
    csv_text = open(rep["manifest"]["csv"], encoding="utf-8").read()
    assert DOI in csv_text
    assert "creativecommons" in csv_text
    assert "reuse-ok-attribution-required" in csv_text


def test_pictures_in_layout_placeholders_are_found(tmp_path):
    """A picture inside a layout placeholder is still a picture.

    python-pptx reports shape_type == PLACEHOLDER (not PICTURE) for these, so a
    shape_type filter drops every image in any deck built on the stock theme
    layouts. Found in the wild: a real 81-slide deck reported 0 of its 61
    pictures.
    """
    prs = Presentation()
    # layout 8 in the default template is "Picture with Caption"
    lay = next(l for l in prs.slide_layouts
               if any(ph.placeholder_format.type == 18 for ph in l.placeholders))
    slide = prs.slides.add_slide(lay)
    ph = next(p for p in slide.placeholders if p.placeholder_format.type == 18)
    img = _make_image(tmp_path / "in-placeholder.png", 77)
    ph.insert_picture(str(img))
    deck = tmp_path / "placeholder-deck.pptx"
    prs.save(str(deck))

    found = [p for s in Presentation(str(deck)).slides for p, _ in iter_pictures(s)]
    assert len(found) == 1, (
        f"picture in a layout placeholder was not found ({len(found)} found); "
        "shape_type reports PLACEHOLDER, not PICTURE")
    assert found[0].image.blob, "placeholder picture exposed no image bytes"
