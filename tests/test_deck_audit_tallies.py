"""The four numbers `figcite audit` reports, tested as rules not as one case.

`deck.audit` produces the headline the whole feature exists for -- "N
substantive but unsourced". Tests existed and asserted
`untagged_substantive == 1` on a fixture where every picture was large and
every record confirmed, so the RULES underneath those numbers were never
exercised. The mutation sweep found all three:

  - `"decorative": w_in < min_inches and h_in < min_inches` survived
    `and -> or`. Under `or`, a picture small in EITHER dimension counts as
    decorative -- so a wide, short figure (a panel strip, a timeline) drops
    out of the unsourced tally entirely and the audit under-reports.
  - `unconfirmed = [... if not r["record"].confirmed]` survived dropping the
    `not`, which inverts the list.
  - `untagged = [... if r["record"] is None and ...]` survived `is -> is not`,
    which also inverts it.

Under-reporting is the direction that matters: this number is what tells you
whether your deck is safe to present.
"""

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from figcite import store
from figcite.deck import audit
from figcite.provenance import Record, embed

import pytest

CITE = "Shiragaki et al. (2020). Phylogenetic Analysis of Capsicum."
DOI = "10.3390/horticulturae6040087"


@pytest.fixture(autouse=True)
def _own_manifest(tmp_path, monkeypatch):
    """One manifest per test.

    conftest gives the whole SESSION a single FIGCITE_HOME, so `store.put`
    calls accumulate across tests in this file -- and `match_picture` falls
    back to a fuzzy dhash search, so a record filed by an earlier test matches
    a deliberately-bare image in a later one and the tally comes out wrong for
    a reason that has nothing to do with the code. Caught exactly that way:
    this file's tests passed one at a time and failed together.
    """
    monkeypatch.setattr(store, "MANIFEST", tmp_path / "manifest.jsonl")
    monkeypatch.setattr(store, "DATA_DIR", tmp_path)


def _image(path, color=7, size=(400, 300)):
    im = Image.new("RGB", size)
    for x in range(size[0]):  # non-flat, so dhash means something
        for y in range(0, size[1], 3):
            im.putpixel((x, y), ((x + color) % 256, (y * 2) % 256, color % 256))
    im.save(path)
    return path


def _tagged(tmp_path, name, *, confirmed=True, color=7):
    src = _image(tmp_path / f"raw-{name}", color)
    rec = Record(
        doi=DOI,
        citation=CITE,
        short_cite="Shiragaki et al. 2020",
        source_kind="pdf-crop",
        confirmed=confirmed,
    )
    out = tmp_path / name
    rec = embed(src, out, rec)
    store.put(rec)
    return out


def _deck(tmp_path, placed, name="deck.pptx"):
    """`placed` is [(image_path, width_in, height_in)] -- sizes are the point
    of this file, so they are explicit rather than defaulted."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    top = 0.3
    for img, w, h in placed:
        slide.shapes.add_picture(
            str(img), Inches(0.4), Inches(top), width=Inches(w), height=Inches(h)
        )
        top += h + 0.2
    p = tmp_path / name
    prs.save(str(p))
    return p


def test_a_wide_short_figure_is_substantive_not_decorative(tmp_path):
    """THE finding.

    6.0 x 0.6 inches: small in one dimension, large in the other -- a panel
    strip or a timeline, which is exactly the kind of figure that needs a
    citation. `and -> or` makes it decorative and it vanishes from the
    unsourced count while the audit still reports success.
    """
    img = _image(tmp_path / "strip.png", color=11, size=(1200, 120))
    deck = _deck(tmp_path, [(img, 6.0, 0.6)])

    rep = audit(str(deck), min_inches=1.0)

    row = rep["rows"][0]
    assert row["decorative"] is False, (
        f"a 6.0x0.6in figure was written off as decorative: {row['size_in']}"
    )
    assert rep["untagged_substantive"] == 1, (
        f"an uncredited 6.0x0.6in figure is missing from the tally: {rep}"
    )


def test_a_figure_small_in_both_dimensions_is_decorative(tmp_path):
    """The positive control. Without it, 'nothing is ever decorative' passes
    the test above while flooding the report with bullet icons."""
    img = _image(tmp_path / "icon.png", color=3, size=(60, 60))
    deck = _deck(tmp_path, [(img, 0.4, 0.4)])

    rep = audit(str(deck), min_inches=1.0)

    assert rep["rows"][0]["decorative"] is True, rep["rows"][0]
    assert rep["untagged_substantive"] == 0, (
        f"a 0.4x0.4in icon was counted as an unsourced figure: {rep}"
    )


def test_the_two_rules_are_counted_together(tmp_path):
    """Both kinds in one deck, so the tally has to discriminate rather than
    apply one verdict to everything."""
    strip = _image(tmp_path / "s.png", color=11, size=(1200, 120))
    icon = _image(tmp_path / "i.png", color=3, size=(60, 60))
    big = _image(tmp_path / "b.png", color=5, size=(400, 300))
    deck = _deck(tmp_path, [(strip, 6.0, 0.6), (icon, 0.4, 0.4), (big, 3.0, 2.2)])

    rep = audit(str(deck), min_inches=1.0)

    assert rep["pictures"] == 3, rep
    assert rep["untagged_substantive"] == 2, (
        f"expected the strip and the big figure, not the icon: {rep}"
    )


def test_an_unconfirmed_record_is_counted_as_unconfirmed(tmp_path):
    """`unconfirmed` survived dropping its `not`, which inverts the list.

    The first version of this test used one confirmed and one unconfirmed
    picture and asserted `unconfirmed == 1` -- which an INVERTED filter also
    returns, because both halves have size 1. It could not fail for the reason
    it named, which is the exact defect this whole audit is hunting, written
    by me while hunting it.

    The counts have to be ASYMMETRIC for a cardinality assertion to
    discriminate a set from its complement.
    """
    yes = _tagged(tmp_path, "confirmed.png", confirmed=True, color=7)
    no1 = _tagged(tmp_path, "unconfirmed-1.png", confirmed=False, color=23)
    no2 = _tagged(tmp_path, "unconfirmed-2.png", confirmed=False, color=61)
    deck = _deck(tmp_path, [(yes, 3.0, 2.2), (no1, 3.0, 2.2), (no2, 3.0, 2.2)])

    rep = audit(str(deck), min_inches=1.0)

    assert rep["tagged"] == 3, f"all three should match the manifest: {rep}"
    assert rep["unconfirmed"] == 2, (
        f"expected the two un-confirmed ones; 1 would mean the filter is inverted: {rep}"
    )


def test_a_tagged_picture_is_not_counted_as_unsourced(tmp_path):
    """`untagged` survived `is -> is not`. Asymmetric for the same reason as
    the test above: 1-and-1 cannot tell a set from its complement."""
    tagged = _tagged(tmp_path, "has-record.png", confirmed=True, color=7)
    bare1 = _image(tmp_path / "no-record-1.png", color=41)
    bare2 = _image(tmp_path / "no-record-2.png", color=83)
    deck = _deck(tmp_path, [(tagged, 3.0, 2.2), (bare1, 3.0, 2.2), (bare2, 3.0, 2.2)])

    rep = audit(str(deck), min_inches=1.0)

    assert rep["tagged"] == 1, rep
    assert rep["untagged_substantive"] == 2, (
        f"expected the two bare ones; 1 would mean the filter is inverted: {rep}"
    )


def test_a_figure_exactly_at_the_threshold_is_substantive(tmp_path):
    """The boundary `<` vs `<=` survived too. At exactly min_inches the
    picture is NOT below the threshold, so it counts -- erring toward
    reporting a figure rather than silently dropping it."""
    img = _image(tmp_path / "exact.png", color=17, size=(300, 300))
    deck = _deck(tmp_path, [(img, 1.0, 1.0)])

    rep = audit(str(deck), min_inches=1.0)

    assert rep["rows"][0]["decorative"] is False, rep["rows"][0]
    assert rep["untagged_substantive"] == 1, (
        f"a figure exactly at the threshold was dropped from the tally: {rep}"
    )


@pytest.mark.parametrize(
    "w_in,h_in",
    [(1.0, 0.5), (0.5, 1.0)],
    ids=["wide-at-threshold", "tall-at-threshold"],
)
def test_only_one_dimension_at_the_threshold_is_still_substantive(tmp_path, w_in, h_in):
    """The predicate is TWO comparisons, and the square fixture above cannot
    tell them apart.

        decorative = w_in < min_inches and h_in < min_inches

    Re-running the logic sweep on the current suite found `Lt -> LtE` still
    surviving here even with the threshold test in place: at 1.0x1.0 both
    comparisons are False, so flipping EITHER to `<=` leaves the `and` False
    and nothing observable changes.

    Asymmetric boundaries separate them -- 1.0x0.5 puts only the first
    comparison on the boundary, 0.5x1.0 only the second -- so each case kills
    exactly one mutant and neither covers for the other. Same rule as the
    rest of this audit: two decisions need two observations.
    """
    img = _image(tmp_path / f"edge-{w_in}x{h_in}.png", color=19, size=(600, 300))
    deck = _deck(tmp_path, [(img, w_in, h_in)])

    rep = audit(str(deck), min_inches=1.0)

    assert rep["rows"][0]["decorative"] is False, (
        f"a {w_in}x{h_in}in figure with one dimension exactly at the threshold "
        f"was written off as decorative: {rep['rows'][0]}"
    )
    assert rep["untagged_substantive"] == 1, rep
