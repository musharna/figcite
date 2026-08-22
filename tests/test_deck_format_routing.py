"""Which implementation handles a file, decided by its suffix -- in two copies.

`_is_pdf` is spelled out twice, identically:

    figcite/cli.py:670       return str(path).lower().endswith(".pdf")
    figcite/service.py:451   return str(path).lower().endswith(".pdf")

and it is not a cosmetic helper. It chooses between two whole
implementations -- `pdfdeck` (PyMuPDF) and `deck` (python-pptx) -- for audit
and for apply, at both the CLI and the service layer.

BOTH copies survived `".pdf" -> "MUTANT"` in the string sweep. Under that
mutation a PDF is handed to python-pptx, which cannot open it. The reason
nothing failed is that every existing PDF test calls `pdfdeck.audit()` /
`pdfdeck.apply()` DIRECTLY. The dispatch that gets you there was never
exercised, so the suite proves the PDF implementation works and proves
nothing about anyone reaching it.

This is the audit's recurring shape: a rule that exists twice needs two
observations, or either copy can be deleted for free. Here it is worse than
usual, because the two copies are in different layers and a fix applied to
one reads as a fix.

Case is asserted too (`.PDF`), since `_is_pdf` lowercases on purpose and a
mutation dropping that would otherwise route a perfectly ordinary
`REPORT.PDF` to the pptx path.
"""

from __future__ import annotations

import random

import fitz
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from figcite import cli, service


def _figure(path, seed, size=(420, 300)):
    rng = random.Random(seed)
    im = Image.new("RGB", size, (252, 250, 248))
    d = ImageDraw.Draw(im)
    for _ in range(24):
        x, y = rng.randrange(size[0] - 50), rng.randrange(size[1] - 50)
        d.rectangle(
            [x, y, x + rng.randrange(15, 60), y + rng.randrange(15, 55)],
            fill=(rng.randrange(256), rng.randrange(256), rng.randrange(256)),
        )
    im.save(path)
    return path


@pytest.fixture
def a_pdf(tmp_path):
    img = _figure(tmp_path / "fig-pdf.png", seed=11)
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(fitz.Rect(60, 60, 400, 300), filename=str(img))
    p = tmp_path / "deck.pdf"
    doc.save(str(p))
    doc.close()
    return p


@pytest.fixture
def a_pptx(tmp_path):
    img = _figure(tmp_path / "fig-pptx.png", seed=23)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(img), Inches(1), Inches(1), Inches(4), Inches(3))
    p = tmp_path / "deck.pptx"
    prs.save(str(p))
    return p


# --- service._is_pdf ------------------------------------------------------


def test_the_service_sends_a_pdf_to_the_pdf_implementation(a_pdf):
    """`kind` and the location wording both come from the branch taken, so
    they say which implementation actually ran -- not merely that something
    produced a report."""
    rep = service.audit(a_pdf)

    assert rep["kind"] == "pdf", rep
    assert rep["rows"], "the PDF produced no rows at all"
    assert rep["rows"][0]["location"].startswith("page "), (
        f"a PDF was audited by the pptx implementation: {rep['rows'][0]}"
    )


def test_the_service_sends_a_pptx_to_the_pptx_implementation(a_pptx):
    """Positive control. Without it, "always route to pdfdeck" satisfies the
    test above while breaking every deck the tool was written for."""
    rep = service.audit(a_pptx)

    assert rep["kind"] == "pptx", rep
    assert rep["rows"], "the pptx produced no rows at all"
    assert rep["rows"][0]["location"].startswith("slide "), (
        f"a pptx was audited by the PDF implementation: {rep['rows'][0]}"
    )


def test_an_uppercase_suffix_is_still_a_pdf(a_pdf, tmp_path):
    """`_is_pdf` lowercases before comparing. Dropping that sends REPORT.PDF
    -- a name a real export produces -- to python-pptx, which cannot open
    it."""
    shouty = tmp_path / "REPORT.PDF"
    shouty.write_bytes(a_pdf.read_bytes())

    rep = service.audit(shouty)

    assert rep["kind"] == "pdf", f"REPORT.PDF was not recognised as a PDF: {rep}"


# --- cli._is_pdf, the second copy ----------------------------------------


def test_the_cli_sends_a_pdf_to_the_pdf_implementation(a_pdf, monkeypatch):
    """The CLI keeps its OWN copy of the predicate, so the service test above
    cannot speak for it. Both implementations are stubbed, so the assertion
    is about which one was CHOSEN rather than about what either produces."""
    called: list[str] = []
    monkeypatch.setattr(
        "figcite.pdfdeck.audit",
        lambda p, **kw: (
            called.append("pdf")
            or {
                "file": str(p),
                "pictures": 0,
                "tagged": 0,
                "unconfirmed": 0,
                "untagged_substantive": 0,
                "rows": [],
            }
        ),
    )
    monkeypatch.setattr(
        "figcite.deck.audit",
        lambda p, **kw: (
            called.append("pptx")
            or {
                "pptx": str(p),
                "pictures": 0,
                "tagged": 0,
                "unconfirmed": 0,
                "untagged_substantive": 0,
                "rows": [],
            }
        ),
    )

    cli.main(["audit", str(a_pdf)])
    assert called == ["pdf"], f"the CLI routed a .pdf to {called}"

    called.clear()
    cli.main(["audit", "deck.pptx"])
    assert called == ["pptx"], f"the CLI routed a .pptx to {called}"


def test_the_two_copies_of_the_predicate_agree(a_pdf):
    """They are separate functions in separate modules with no shared source.

    Nothing makes them agree except that someone typed the same thing twice,
    so this pins the property directly: fixing one and not the other is the
    failure this file exists to catch.
    """
    for name in ("deck.pdf", "REPORT.PDF", "deck.pptx", "notes.txt", "archive.pdf.bak"):
        assert cli._is_pdf(name) == service._is_pdf(name), (
            f"cli._is_pdf and service._is_pdf disagree about {name!r}: "
            f"{cli._is_pdf(name)} vs {service._is_pdf(name)}"
        )


def test_a_name_merely_containing_pdf_is_not_a_pdf():
    """`endswith`, not `in`. `archive.pdf.bak` is a backup file; routing it to
    PyMuPDF because the name has 'pdf' in it somewhere is the open-set error
    this suite keeps finding."""
    assert not service._is_pdf("archive.pdf.bak")
    assert not service._is_pdf("pdf_notes.pptx")
    assert service._is_pdf("real.pdf")
