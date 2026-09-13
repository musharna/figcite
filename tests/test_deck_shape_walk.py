"""A deck holds shapes that are not pictures, and none of them were ever in one.

`deck.iter_pictures` walks a slide's shape tree by XML tag:

    if   t == "grpSp":  recurse into the group
    elif t == "pic":    this is a picture

Both comparisons survived the reduced-ROR sweep widened to `<=`, and both
survive for the same reason: every deck any test has ever built contains
pictures, groups, and nothing else. The comparison is against a NAME, and
nothing in the fixtures sorts near enough to that name to tell a widened one
apart.

Two ordinary shapes do. A table or chart is a `graphicFrame` and a connector
is a `cxnSp`, and

    "graphicFrame" < "grpSp"     ('a' < 'p' at index 2)
    "graphicFrame" < "pic"       ('g' < 'p')
    "cxnSp"        < "pic"       ('c' < 'p')

so widening either branch pulls them in:

  `t <= "pic"`    yields the table AS A PICTURE. figcite then treats a table
                  as a figure -- looks for its provenance, counts it in the
                  deck tally, captions it.

  `t <= "grpSp"`  treats the table as a GROUP and recurses into `sh.shapes`,
                  which a graphicFrame does not have.

A deck with a table exercises both. Decks have tables.
"""

from __future__ import annotations

import pytest

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches
from PIL import Image

from figcite.deck import iter_pictures


def _png(path, colour=(30, 90, 140)):
    Image.new("RGB", (40, 30), colour).save(path)
    return path


@pytest.fixture
def deck_with_a_table(tmp_path):
    """One picture, one table, one connector, on a single slide.

    The picture is the positive control: without it, "found nothing" would
    satisfy every assertion below.
    """
    img = _png(tmp_path / "fig.png")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(img), Inches(0.5), Inches(0.5), width=Inches(2))
    slide.shapes.add_table(2, 2, Inches(3), Inches(0.5), Inches(3), Inches(1))
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(4), Inches(3), Inches(4))
    path = tmp_path / "mixed.pptx"
    prs.save(str(path))
    return path


def _tags(path):
    from figcite.deck import _tag

    prs = Presentation(str(path))
    return sorted(_tag(sh) for sh in prs.slides[0].shapes)


def test_the_fixture_really_contains_the_shapes_that_sort_below(deck_with_a_table):
    """The premise. If python-pptx ever stopped emitting `graphicFrame` for a
    table, the tests below would keep passing on a deck that no longer holds
    anything a widened comparison could pull in."""
    tags = _tags(deck_with_a_table)
    assert "pic" in tags, tags
    assert "graphicFrame" in tags, f"no table in the fixture: {tags}"
    assert any(t < "pic" for t in tags if t != "pic"), (
        f"nothing in this deck sorts below 'pic', so it cannot exercise the widening: {tags}"
    )
    assert any(t < "grpSp" for t in tags), f"nothing in this deck sorts below 'grpSp': {tags}"


def test_a_table_is_not_a_picture(deck_with_a_table):
    """`t == "pic"` widened to `<=` yields the table and the connector too."""
    prs = Presentation(str(deck_with_a_table))
    found = list(iter_pictures(prs.slides[0]))

    assert len(found) == 1, (
        f"expected exactly the one picture; a non-picture shape was collected "
        f"as a figure: {[sh.shape_type for sh, _ in found]}"
    )
    from figcite.deck import _tag

    assert _tag(found[0][0]) == "pic", _tag(found[0][0])


def test_walking_a_deck_with_a_table_does_not_explode(deck_with_a_table):
    """`t == "grpSp"` widened to `<=` recurses into a graphicFrame's `.shapes`,
    which does not exist. Asserted separately from the count above because it
    fails as an AttributeError rather than a wrong answer, and a test that only
    counts would report that as an error rather than a finding."""
    prs = Presentation(str(deck_with_a_table))
    found = list(iter_pictures(prs.slides[0]))  # must not raise
    assert found, "the picture was lost"


def test_a_picture_inside_a_group_is_still_found(tmp_path):
    """Positive control for the `grpSp` branch itself. Every test above would
    pass on a walker that never recursed at all."""
    img = _png(tmp_path / "grouped.png", (200, 60, 20))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = [
        slide.shapes.add_picture(str(img), Inches(0.5), Inches(0.5), width=Inches(1)),
        slide.shapes.add_picture(str(img), Inches(2.0), Inches(0.5), width=Inches(1)),
    ]
    try:
        slide.shapes.add_group_shape(shapes)
    except (AttributeError, TypeError) as e:  # pragma: no cover - version-dependent
        pytest.skip(f"this python-pptx cannot build a group shape: {e}")
    path = tmp_path / "grouped.pptx"
    prs.save(str(path))

    prs2 = Presentation(str(path))
    found = list(iter_pictures(prs2.slides[0]))
    assert len(found) == 2, f"pictures inside a group were not walked into: {len(found)}"
