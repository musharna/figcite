"""A record read out of an image must identify the bytes it was read from.

`embed` serialises the record into the PNG's tEXt chunk and only THEN hashes
the finished file, so the copy inside the image can never carry its own
sha256 -- the value did not exist yet. `read_embedded` handed that copy back
verbatim, sha256 == "", and the deck screen built its thumbnail ref from it:
`sha:` with nothing after the colon, a 404, and "(image unavailable)" on the
one match path a tagged deck hits most.

The JPEG branch of `read_embedded` already stamps the blob's hashes onto the
recovered record. The PNG branch is the same situation and gets the same
treatment: the reader knows the bytes, so the record it returns names them.
"""

from PIL import Image

from figcite import service, store
from figcite.provenance import Record, dhash_bytes, embed, read_embedded, sha256_bytes


def _tagged_png(tmp_path):
    src = tmp_path / "fig.png"
    im = Image.new("RGB", (64, 48))
    px = im.load()
    for x in range(64):
        for y in range(48):
            px[x, y] = ((x * 37 + y * 101) % 256, (x * 7) % 256, (y * 11) % 256)
    im.save(src)
    dst = tmp_path / "fig.tagged.png"
    rec = Record(doi="10.1/x", citation="Someone 2020", confirmed=True)
    written = embed(src, dst, rec)
    return dst, written


def test_a_record_read_from_a_png_carries_that_pngs_hashes(tmp_path):
    dst, written = _tagged_png(tmp_path)
    blob = dst.read_bytes()

    back = read_embedded(blob)

    assert back is not None and back.doi == "10.1/x", "positive control: the record is there"
    assert back.sha256 == sha256_bytes(blob), (
        f"read_embedded returned sha256={back.sha256!r} for a blob whose sha256 is "
        f"{sha256_bytes(blob)[:12]}...: the record does not identify its own bytes"
    )
    assert back.dhash == dhash_bytes(blob)
    # And it agrees with what `embed` recorded in the manifest/sidecar.
    assert back.sha256 == written.sha256


def test_the_deck_screen_gets_a_thumbnail_ref_for_an_embedded_match(tmp_path, monkeypatch):
    """End to end at the layer the symptom appeared: the audit row's `ref`."""
    from pptx import Presentation
    from pptx.util import Inches

    dst, written = _tagged_png(tmp_path)
    monkeypatch.setattr(store, "all_records", lambda: {written.sha256: written})
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(dst), Inches(1), Inches(1), width=Inches(4))
    deckpath = tmp_path / "d.pptx"
    prs.save(deckpath)

    rows = service.audit(str(deckpath))["rows"]

    assert rows and rows[0]["matched_by"] == "embedded-metadata", rows
    assert rows[0]["ref"] == f"sha:{written.sha256}", (
        f"thumbnail ref is {rows[0]['ref']!r}; an empty sha is a 404 and '(image unavailable)'"
    )
