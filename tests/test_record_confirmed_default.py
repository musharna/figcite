"""`confirmed` defaults to False, and a manifest that omits it must not lie.

    confirmed: bool = False  # False => a machine guessed this; never auto-cite it

That comment states the tool's central rule, and the mutation sweep flipped
the default to True without a single test failing.

Triage note, because it changes the severity: this is NOT reachable from
normal construction. Every `Record(...)` call in the package passes
`confirmed` explicitly, which is exactly why the mutant survived. The default
is reached on ONE path -- `Record.from_dict`, which filters an incoming dict
to known fields and calls `cls(**subset)`. A manifest line written by an older
version, hand-edited, or truncated therefore takes the default.

That path is the one reading persisted data of unknown origin, so "a record
whose confirmed flag is missing is treated as CONFIRMED" is the failure worth
pinning: unknown provenance must not default to trusted.
"""

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from figcite import store
from figcite.deck import apply as deck_apply
from figcite.provenance import Record, embed


def test_a_bare_record_is_not_confirmed():
    assert Record().confirmed is False


def test_a_manifest_line_without_the_field_is_not_confirmed():
    """THE reachable case. `from_dict` drops unknown keys and relies on the
    dataclass defaults for absent ones, so an older or partial manifest row
    lands here."""
    rec = Record.from_dict({"sha256": "a" * 64, "doi": "10.1/x", "citation": "Someone et al. 2020"})

    assert rec.confirmed is False, (
        "a manifest row with no `confirmed` field deserialised as confirmed, "
        "so a record of unknown provenance would be auto-cited"
    )


def test_an_explicit_true_still_survives_the_round_trip():
    """Positive control: if `from_dict` ignored the field entirely, the test
    above would pass while confirmation could never be persisted at all."""
    rec = Record.from_dict({"sha256": "b" * 64, "confirmed": True})
    assert rec.confirmed is True


def _image(path, color=7):
    im = Image.new("RGB", (400, 300))
    for x in range(400):
        for y in range(0, 300, 3):
            im.putpixel((x, y), ((x + color) % 256, (y * 2) % 256, color % 256))
    im.save(path)
    return path


def test_an_unconfirmed_record_does_not_reach_a_slide(tmp_path, monkeypatch):
    """The consequence, asserted end to end.

    A test on the flag alone would keep passing if the deck writer stopped
    consulting it. `apply` must caption a confirmed record and must NOT
    caption an unconfirmed one unless explicitly overridden.
    """
    monkeypatch.setattr(store, "MANIFEST", tmp_path / "manifest.jsonl")
    monkeypatch.setattr(store, "DATA_DIR", tmp_path)

    src = _image(tmp_path / "raw.png", color=9)
    rec = Record(
        doi="10.1/guess",
        citation="A Guessed Citation 2020",
        short_cite="Guess 2020",
        source_kind="clipboard",
        confirmed=False,
    )
    filed = embed(src, tmp_path / "fig.png", rec)
    store.put(filed)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(tmp_path / "fig.png"), Inches(1), Inches(1), width=Inches(3))
    deck_path = tmp_path / "deck.pptx"
    prs.save(str(deck_path))

    out = tmp_path / "deck.cited.pptx"
    deck_apply(str(deck_path), str(out))

    text = _all_text(out)
    assert "A Guessed Citation 2020" not in text, (
        "an unconfirmed citation reached the slide without an override"
    )

    out2 = tmp_path / "deck.forced.pptx"
    deck_apply(str(deck_path), str(out2), allow_unconfirmed=True)
    assert "A Guessed Citation 2020" in _all_text(out2), (
        "the explicit override did not work, so the test above proves nothing"
    )


def _all_text(pptx_path) -> str:
    prs = Presentation(str(pptx_path))
    bits = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                bits.append(shape.text_frame.text)
    return "\n".join(bits)
