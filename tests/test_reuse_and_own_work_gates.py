"""Three gates about what a figure's provenance ENTITLES you to say.

All three survived the string sweep, and two of them are the same rule
written twice:

  bibtex.py:110   if rec.reuse and rec.reuse != "unknown":
  deck.py:326     elif rec.reuse and rec.reuse != "unknown":
  pdfdeck.py:185  own_work = rec.source_kind == "generated"

`"unknown"` is the sentinel `reuse` carries when the licence could not be
determined. The check exists so an undetermined licence is not printed as if
it were a finding. Mutate the sentinel and the guard stops recognising it:
every figure whose licence is unknown starts announcing `Reuse: unknown` on
the credits page and `reuse: unknown` in the .bib -- a statement about
rights, derived from nothing, on a document that gets forwarded.

The third is `own_work`, and the asymmetry is the point. The pptx copy
(`deck.py:346`) is KILLED by `test_deck.py::test_own_work_is_not_captioned_by_default`;
the PDF copy at `pdfdeck.py:185` survives. Same rule, one twin covered and
one not -- the fourth time this audit has found exactly that shape, after
the audit tallies, the decorative predicate, and the confirmed gate.
"""

from __future__ import annotations

import random

import fitz
import pytest
from PIL import Image, ImageDraw

from pptx import Presentation
from pptx.util import Inches

from figcite import store
from figcite.bibtex import records_to_bibtex
from figcite.deck import apply as pptx_apply
from figcite.pdfdeck import apply as pdf_apply
from figcite.provenance import Record, embed

CITE = "Shiragaki et al. (2020). Horticulturae 6: 87."
DOI = "10.3390/horticulturae6040087"


@pytest.fixture(autouse=True)
def _own_manifest(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "MANIFEST", tmp_path / "manifest.jsonl")
    monkeypatch.setattr(store, "DATA_DIR", tmp_path)


def _rec(**kw):
    base = dict(
        doi=DOI,
        citation=CITE,
        short_cite="Shiragaki et al. 2020",
        confirmed=True,
        source_kind="pdf-crop",
    )
    base.update(kw)
    return Record(**base)


# --- the reuse sentinel, in bibtex -----------------------------------------


def test_an_unknown_licence_is_not_written_into_the_bib():
    """THE finding. `unknown` means "could not determine", and a .bib note
    reading `reuse: unknown` presents that as a determination."""
    out = records_to_bibtex([_rec(reuse="unknown")])

    assert "reuse: unknown" not in out["bibtex"], (
        f"an undetermined licence was written into the .bib as a note:\n{out['bibtex']}"
    )


def test_a_real_licence_is_written_into_the_bib():
    """Positive control. "Never write reuse" passes the test above while
    losing the licence on every figure that actually has one."""
    out = records_to_bibtex([_rec(reuse="CC BY 4.0")])

    assert "reuse: CC BY 4.0" in out["bibtex"], (
        f"a known licence never reached the .bib:\n{out['bibtex']}"
    )


def test_an_empty_reuse_is_also_withheld():
    """The `rec.reuse and ...` half of the same expression."""
    out = records_to_bibtex([_rec(reuse="")])
    assert "reuse:" not in out["bibtex"], out["bibtex"]


# --- the reuse sentinel, in the deck credits (the second copy) -------------


def _figure(path, seed, size=(420, 300)):
    rng = random.Random(seed)
    im = Image.new("RGB", size, (250, 250, 248))
    d = ImageDraw.Draw(im)
    for _ in range(24):
        x, y = rng.randrange(size[0] - 60), rng.randrange(size[1] - 60)
        d.rectangle(
            [x, y, x + rng.randrange(15, 60), y + rng.randrange(15, 55)],
            fill=(rng.randrange(256), rng.randrange(256), rng.randrange(256)),
        )
    im.save(path)
    return path


def _staged_pdf(tmp_path, rec, name="export.pdf", seed=7):
    raw = _figure(tmp_path / f"raw-{seed}.png", seed=seed)
    img = tmp_path / f"fig-{seed}.png"
    rec = embed(raw, img, rec)
    store.put(rec)

    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(fitz.Rect(60, 60, 400, 290), filename=str(img))
    p = tmp_path / name
    doc.save(str(p))
    doc.close()
    return p


def _credits_text(out_path):
    doc = fitz.open(str(out_path))
    text = doc[doc.page_count - 1].get_text()
    doc.close()
    return text


def _staged_pptx(tmp_path, rec, name="deck.pptx", seed=51):
    """The `Reuse:` fallback is a PPTX-ONLY feature, so this drives pptx.

    The first version of the two tests below drove `pdfdeck.apply`, and
    `pdfdeck._credit_line` (pdfdeck.py:236) prints a licence ONLY when
    `license_url` is set -- it has no `elif rec.reuse and rec.reuse !=
    "unknown"` fallback at all. So "unknown is not printed" held for a reason
    with nothing to do with the guard: on that path reuse is never printed
    under any value. The test could not have failed.

    The positive control below is the only thing that caught it.

    (The divergence itself is real: a figure with a known `reuse` but no
    `license_url` gets a credit line in a pptx deck and none in a PDF. That
    reads as a deliberate difference rather than a defect -- noted, not
    changed.)
    """
    raw = _figure(tmp_path / f"raw-{seed}.png", seed=seed)
    img = tmp_path / f"fig-{seed}.png"
    rec = embed(raw, img, rec)
    store.put(rec)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(img), Inches(0.5), Inches(0.5), width=Inches(3))
    p = tmp_path / name
    prs.save(str(p))
    return p


def _pptx_text(path):
    prs = Presentation(str(path))
    out = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
    return "\n".join(out)


def test_an_unknown_licence_is_not_printed_on_the_credits_page(tmp_path):
    """The deck copy of the same sentinel. Asserted separately from the
    bibtex one because they are separate expressions -- fixing one and not
    the other is the failure this file is built to catch."""
    deck = _staged_pptx(tmp_path, _rec(reuse="unknown"), name="unknown.pptx", seed=31)
    out = tmp_path / "unknown.cited.pptx"

    pptx_apply(deck, out, captions=True, credits=True)

    assert "Reuse: unknown" not in _pptx_text(out), (
        f"the credits page asserted a licence it does not know:\n{_pptx_text(out)}"
    )


def test_a_real_licence_is_printed_on_the_credits_page(tmp_path):
    """Positive control -- and the one that caught this file testing the
    wrong module. If a known licence does not appear here, the "unknown is
    withheld" test above is passing because NOTHING is ever printed."""
    deck = _staged_pptx(tmp_path, _rec(reuse="CC BY 4.0"), name="cc.pptx", seed=37)
    out = tmp_path / "cc.cited.pptx"

    pptx_apply(deck, out, captions=True, credits=True)

    assert "CC BY 4.0" in _pptx_text(out), (
        f"a known licence never reached the credits page:\n{_pptx_text(out)}"
    )


# --- own work, in the PDF twin --------------------------------------------


def test_a_generated_figure_is_not_captioned_by_default_in_a_pdf(tmp_path):
    """THE PDF-twin finding.

    `test_deck.py::test_own_work_is_not_captioned_by_default` pins this for
    pptx and kills the pptx mutant. The PDF path spells the rule out again
    at pdfdeck.py:185 and nothing pinned it, so `source_kind == "generated"`
    could be changed freely.

    Your own plot does not need someone else's citation stamped across it.

    ASSERT ON WHAT THE CAPTION ACTUALLY RENDERS. The first version of this
    test set `citation="My own figure"` and asserted that string was absent
    -- but `_caption_text` (pdfdeck.py:228) renders

        f"[{n}] {rec.short_cite or rec.citation[:60]}"

    and `short_cite` was still the fixture default, so the caption would have
    read "[1] Shiragaki et al. 2020". The asserted string appeared in neither
    the captioned nor the uncaptioned case, so the test passed under the
    mutation too. It was caught by re-running the mutant, not by reading it:
    the suite went green and the mutant lived.
    """
    rec = _rec(
        source_kind="generated",
        doi="",
        citation="A plot I made myself",
        short_cite="MY-OWN-WORK-MARKER",
    )
    pdf = _staged_pdf(tmp_path, rec, name="mine.pdf", seed=41)
    out = tmp_path / "mine.cited.pdf"

    pdf_apply(pdf, out, captions=True, credits=True)

    doc = fitz.open(str(out))
    figure_page = doc[0].get_text()
    doc.close()

    assert "MY-OWN-WORK-MARKER" not in figure_page, (
        f"own work was captioned as if it were someone else's: {figure_page!r}"
    )


def test_a_sourced_figure_is_still_captioned_in_a_pdf(tmp_path):
    """Positive control. "Caption nothing" passes the test above and makes
    the entire feature a no-op."""
    pdf = _staged_pdf(
        tmp_path, _rec(source_kind="pdf-crop"), name="theirs.pdf", seed=43
    )
    out = tmp_path / "theirs.cited.pdf"

    pdf_apply(pdf, out, captions=True, credits=True)

    doc = fitz.open(str(out))
    figure_page = doc[0].get_text()
    doc.close()

    assert "Shiragaki" in figure_page, (
        f"a sourced figure lost its caption: {figure_page!r}"
    )


# --- own_work, from the side of the comparison nothing stood on -------------
#
# `source_kind == "generated"` means "the user made this themselves", and it
# buys two exemptions: bibtex drops the record from the bibliography, and both
# deck writers suppress its citation caption. Widened to `<=` it stops naming a
# kind and starts naming a HALF-LINE -- and the kinds that fall inside it are
# `clipboard`, `clipboard-from-pdf` and `clipboard-from-web`, which is figcite's
# primary capture path and by definition somebody ELSE's figure.
#
# So the mutant's effect is: a figure snipped out of a paper is reclassified as
# your own work, dropped from the bibliography, and printed with no attribution.
# That is the exact outcome this tool exists to prevent.
#
# The tests above already cover both sides of `own_work` -- and could not see
# it. `_rec` defaults to `source_kind="pdf-crop"`, and "pdf-crop" > "generated",
# so the sourced-figure control sits on the one side of the comparison the
# widening does not reach. It proves the feature works; it cannot prove the
# predicate discriminates.

CAPTURED_KINDS = ["clipboard", "clipboard-from-pdf", "clipboard-from-web"]


def test_the_captured_kinds_really_do_sort_below_generated():
    """The premise the three tests below rest on.

    If a kind stopped sorting below "generated" it would stop exercising the
    widening, and those tests would keep passing while testing nothing -- which
    is precisely how `pdf-crop` came to be the sourced representative.
    """
    for kind in CAPTURED_KINDS:
        assert kind < "generated", f"{kind!r} no longer sorts below 'generated'"
    assert "pdf-crop" > "generated", (
        "pdf-crop now sorts below 'generated'; the older tests in this file "
        "would start covering the widening and this comment would be wrong"
    )


@pytest.mark.parametrize("kind", CAPTURED_KINDS)
def test_a_captured_figure_is_not_counted_as_own_work_in_the_bib(kind):
    """Someone else's figure must appear in the bibliography."""
    rep = records_to_bibtex([_rec(source_kind=kind)])
    assert rep["skipped_own_work"] == 0, (
        f"a {kind!r} capture was counted as the user's own work and dropped "
        f"from the bibliography: {rep}"
    )
    assert rep["included"] == 1, rep
    # The DOI, not the short cite: `_rec` sets no authors/year, so the entry
    # renders as @article{anonnd} and carries the DOI as its only identifying
    # field. Asserting on a name the entry never contains would fail for a
    # reason that has nothing to do with the own-work gate.
    assert DOI in rep["bibtex"], rep["bibtex"]


def test_a_captured_figure_is_still_captioned_in_a_pdf(tmp_path):
    """One representative rather than all three: the PDF and pptx writers are
    slow, and `clipboard` is both the primary path and the lexicographic floor
    of the widened range, so it is the strongest single case."""
    pdf = _staged_pdf(
        tmp_path, _rec(source_kind="clipboard"), name="captured.pdf", seed=44
    )
    out = tmp_path / "captured.cited.pdf"

    pdf_apply(pdf, out, captions=True, credits=True)

    doc = fitz.open(str(out))
    figure_page = doc[0].get_text()
    doc.close()
    assert "Shiragaki" in figure_page, (
        f"a captured figure was treated as own work and lost its caption: "
        f"{figure_page!r}"
    )


def _first_slide_text(path):
    """Text on the FIGURE slide only.

    `_pptx_text` concatenates every slide, credits page included -- and the
    credits page prints the citation whether or not the caption was
    suppressed. So a whole-deck search cannot tell "captioned" from
    "credited", and the first version of the test below passed while its
    mutant survived. The PDF test above already scoped itself to `doc[0]`
    for exactly this reason.
    """
    prs = Presentation(str(path))
    slide = prs.slides[0]
    return "\n".join(
        sh.text_frame.text for sh in slide.shapes if sh.has_text_frame
    )


def test_a_captured_figure_is_still_captioned_in_a_pptx(tmp_path):
    deck = _staged_pptx(
        tmp_path, _rec(source_kind="clipboard"), name="captured.pptx", seed=52
    )
    out = tmp_path / "captured.cited.pptx"

    pptx_apply(deck, out, captions=True, credits=True)

    figure_slide = _first_slide_text(out)
    assert "Shiragaki" in figure_slide, (
        f"a captured figure was treated as own work and lost its caption: "
        f"{figure_slide!r}"
    )


def test_the_pptx_probe_reads_the_figure_slide_not_the_credits(tmp_path):
    """Positive control for `_first_slide_text`.

    A helper that returned the whole deck would make the test above pass on a
    suppressed caption, because the credits page names the work anyway. This
    pins that the figure slide and the credits page are different text.
    """
    deck = _staged_pptx(
        tmp_path, _rec(source_kind="generated"), name="own.pptx", seed=53
    )
    out = tmp_path / "own.cited.pptx"

    pptx_apply(deck, out, captions=True, credits=True)

    assert "Shiragaki" not in _first_slide_text(out), (
        "own work was captioned on the figure slide, or the probe is reading "
        "the whole deck"
    )
    assert "Shiragaki" in _pptx_text(out), (
        "the credits page stopped naming the work, so the contrast this "
        "control draws no longer exists"
    )


def test_no_source_kind_sorts_below_clipboard():
    """What makes `autostart.status`'s widened count harmless -- checked.

    `status` counts captures with `r.source_kind == "clipboard"`. Widened to
    `<=` that count would absorb every kind sorting below "clipboard", and
    today there are none: "clipboard" is a prefix of "clipboard-from-pdf" and
    "clipboard-from-web", so it is the floor.

    That is an argument about VALUES, not about the code path -- `_actions.py`
    takes `kind: str` and hands it straight to `Record(source_kind=kind)`, so
    nothing structurally stops a caller inventing a new one. Which makes it a
    hypothesis, and the honest thing to do with a hypothesis is give it
    something that can falsify it. Add a kind sorting below "clipboard" and
    this fails, pointing at the equivalence claim that would go stale.
    """
    import ast
    import pathlib

    import figcite

    literals: dict[str, str] = {}
    for path in sorted(pathlib.Path(figcite.__file__).parent.glob("*.py")):
        tree = ast.parse(path.read_text())
        for node in ast.walk(tree):
            if not isinstance(node, ast.Call):
                continue
            for kw in node.keywords:
                if kw.arg in ("source_kind", "kind") and isinstance(
                    kw.value, ast.Constant
                ):
                    if isinstance(kw.value.value, str):
                        literals[kw.value.value] = f"{path.name}:{node.lineno}"

    assert literals, "found no source_kind literals at all; this probe is broken"
    below = {k: where for k, where in literals.items() if k < "clipboard"}
    assert not below, (
        f"a source_kind now sorts below 'clipboard': {below}. "
        f"`autostart.status`'s count is no longer safe under a widened "
        f"comparison, and the equivalence claim for it is stale."
    )
