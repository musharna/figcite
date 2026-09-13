import json
import socket
import threading
import time
import urllib.error
import urllib.request

import pytest
from PIL import Image
from pptx import Presentation

from figcite import crossref, service, web


def test_the_server_serves_pending_over_a_real_socket():
    srv = web.make_server(0)  # ephemeral port
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{port}/api/pending") as r:
            assert r.status == 200
            assert isinstance(json.loads(r.read())["items"], list)
    finally:
        srv.shutdown()


def test_it_binds_loopback_only():
    srv = web.make_server(0)
    try:
        assert srv.server_address[0] == "127.0.0.1"
    finally:
        srv.server_close()


def test_it_refuses_a_port_already_in_use_rather_than_moving():
    held = socket.socket()
    held.bind(("127.0.0.1", 0))
    held.listen(1)
    port = held.getsockname()[1]
    try:
        with pytest.raises(OSError):
            web.make_server(port)
    finally:
        held.close()


# --- round 1 security fix: /api/confirm's splat exposed confirm()'s whole
# signature to the wire, including `out` (an arbitrary file-write path) ---


def _stage(tmp_path, monkeypatch, name="clip-1"):
    """Mirrors tests/test_service_confirm.py's `_stage` helper."""
    staging = tmp_path / "staging"
    staging.mkdir(exist_ok=True)
    png = staging / f"{name}.png"
    Image.new("RGB", (4, 4), (30, 90, 140)).save(png)
    (staging / f"{name}.pending.json").write_text(
        json.dumps(
            {
                "png": str(png),
                "capture": {"process": "firefox", "title": "A paper"},
                "inference": {},
            }
        )
    )
    monkeypatch.setattr(service.clipboard, "staging_dirs", lambda: (None, staging))
    return png


def _post(port, path, payload, headers=None, timeout=None):
    data = json.dumps(payload).encode()
    req = urllib.request.Request(
        f"http://127.0.0.1:{port}{path}",
        data=data,
        headers={"Content-Type": "application/json", **(headers or {})},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, json.loads(r.read())
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read())


def _raw_request(port, raw_bytes: bytes, timeout: float = 5) -> bytes:
    """Send exactly `raw_bytes` over a fresh socket and return the full
    response. For wire-level probes (bad Content-Length, forged Host) that
    urllib's Request can't express, because it computes its own headers.
    """
    with socket.create_connection(("127.0.0.1", port), timeout=timeout) as s:
        s.sendall(raw_bytes)
        s.settimeout(timeout)
        chunks = []
        try:
            while True:
                chunk = s.recv(4096)
                if not chunk:
                    break
                chunks.append(chunk)
        except socket.timeout:
            pytest.fail(f"no response within {timeout}s -- the handler thread hung")
    return b"".join(chunks)


def _status_of(raw_response: bytes) -> bytes:
    return raw_response.split(b"\r\n", 1)[0]


def test_confirm_out_field_cannot_write_outside_the_library(tmp_path, monkeypatch):
    """The regression test for the arbitrary-file-write finding.

    `service.confirm(ref, out=...)` is a real CLI capability
    (`figcite confirm -o/--out`), but `/api/confirm` splatting `**payload`
    straight into `confirm()` made the wire contract equal to the whole
    function signature, so any client could set `out=` to any path this
    process can write and have PNG bytes land there. Must fail against the
    committed, un-allowlisted route: the victim file's bytes change.

    Round 2 (M1) flipped the allowlist from silently dropping an unknown
    field to rejecting the whole request with a 400 -- so `out` is now
    refused outright rather than quietly ignored. The byte-level assertion
    stays: it is the property that actually matters, independent of which
    status code got there.
    """
    _stage(tmp_path, monkeypatch, name="clip-write-atk")
    victim = tmp_path / "VICTIM.txt"
    victim.write_bytes(b"do not touch")

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        status, body = _post(
            port,
            "/api/confirm",
            {
                "ref": "staged:clip-write-atk.png",
                "cite": "Attacker 2026",
                "out": str(victim),
            },
        )
        assert status == 400, body
        assert "out" in body.get("error", ""), body
        assert victim.read_bytes() == b"do not touch", (
            "the victim file's bytes changed -- `out` reached finalize()"
        )
    finally:
        srv.shutdown()


def test_confirm_rejects_an_unknown_field_instead_of_silently_dropping_it(tmp_path, monkeypatch):
    """Round 2, M1. A silently-dropped unknown field means the allowlist can
    never report being too narrow: if the UI later sends a field this route
    doesn't recognise yet, the server would answer 200 and quietly ignore it.
    Rejecting with 400 makes that gap fail loud instead -- same security
    property (the field never reaches `confirm()`), louder failure mode.
    """
    _stage(tmp_path, monkeypatch, name="clip-junk")

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        status, body = _post(
            port,
            "/api/confirm",
            {"ref": "staged:clip-junk.png", "cite": "Junk 2026", "nonsense": 1},
        )
        assert status == 400, body
        assert "TypeError" not in json.dumps(body)
        assert "nonsense" in body.get("error", ""), body
    finally:
        srv.shutdown()


def test_confirm_still_files_a_real_capture_to_the_library(tmp_path, monkeypatch):
    """Positive control: a normal confirm through the API still works."""
    _stage(tmp_path, monkeypatch, name="clip-positive")

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        status, body = _post(
            port,
            "/api/confirm",
            {"ref": "staged:clip-positive.png", "cite": "Positive Control 2026"},
        )
        assert status == 200, body
        assert body["ok"] is True
        assert body["citation"] == "Positive Control 2026"

        with urllib.request.urlopen(f"http://127.0.0.1:{port}/api/pending") as r:
            refs = [i["ref"] for i in json.loads(r.read())["items"]]
        assert "staged:clip-positive.png" not in refs
    finally:
        srv.shutdown()

    matches = [
        r
        for r in service.store.all_records().values()
        if r.citation == "Positive Control 2026" and r.confirmed
    ]
    assert matches, "no confirmed library record was filed"


def test_apply_route_never_forwards_allow_unconfirmed(monkeypatch):
    """No path from this API to `allow_unconfirmed`, verified, not just read.

    Round 2 (M1): the allowlist now rejects an unrecognised field with 400
    instead of silently dropping it, so this doesn't just fail to forward
    `allow_unconfirmed` -- the request never reaches `service.apply` at all.
    `calls == []` proves the stub was never called.
    """
    calls = []

    def _stub_apply(path, out=None, **opts):
        calls.append(opts)
        return {"ok": True}

    monkeypatch.setattr(web.service, "apply", _stub_apply)

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        status, body = _post(
            port, "/api/apply", {"path": "irrelevant.pptx", "allow_unconfirmed": True}
        )
        assert status == 400, body
        assert "allow_unconfirmed" in body.get("error", ""), body
    finally:
        srv.shutdown()

    assert calls == [], f"allow_unconfirmed reached service.apply's opts: {calls!r}"


def test_it_refuses_a_cross_origin_post(tmp_path, monkeypatch):
    _stage(tmp_path, monkeypatch, name="clip-csrf-evil")

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        status, body = _post(
            port,
            "/api/skip",
            {"ref": "staged:clip-csrf-evil.png"},
            headers={"Origin": "http://evil.example"},
        )
        assert status == 403, body
    finally:
        srv.shutdown()


def test_same_origin_and_no_origin_posts_still_work(tmp_path, monkeypatch):
    _stage(tmp_path, monkeypatch, name="clip-csrf-ok1")

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        # no Origin header at all (curl, or a same-page fetch without CORS mode)
        status, body = _post(port, "/api/skip", {"ref": "staged:clip-csrf-ok1.png"})
        assert status == 200, body

        # Origin matching the server's own address
        status, body = _post(
            port,
            "/api/skip",
            {"ref": "irrelevant"},
            headers={"Origin": f"http://127.0.0.1:{port}"},
        )
        assert status == 200, body
    finally:
        srv.shutdown()


# --- round 2 security fixes ---------------------------------------------


def test_malformed_json_body_gets_a_clean_400_not_a_dropped_connection():
    """I1. Body parsing used to sit OUTSIDE the try block, so a malformed
    request got no HTTP response at all -- a dropped connection, which is
    strictly worse than the bare 500 this handler exists to prevent.
    `json.JSONDecodeError` is a `ValueError`; it becomes a clean 400 once
    reading and parsing the body happen inside the try.
    """
    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        req = urllib.request.Request(
            f"http://127.0.0.1:{port}/api/skip",
            data=b"{not json",
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with pytest.raises(urllib.error.HTTPError) as ei:
            urllib.request.urlopen(req)
        assert ei.value.code == 400
    finally:
        srv.shutdown()


def test_negative_content_length_gets_400_not_a_hang():
    """I2. `Content-Length: -1` makes stdlib's `self.rfile.read(-1)` read
    until EOF, blocking the handler thread as long as the client holds the
    connection open -- unbounded, on a `ThreadingHTTPServer`. Must be
    rejected before `read()` is ever called, not after.
    """
    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        resp = _raw_request(
            port,
            b"POST /api/skip HTTP/1.1\r\n"
            + f"Host: 127.0.0.1:{port}\r\n".encode()
            + b"Content-Type: application/json\r\n"
            b"Content-Length: -1\r\n"
            b"Connection: close\r\n"
            b"\r\n",
            timeout=5,
        )
        assert b" 400 " in _status_of(resp), resp
    finally:
        srv.shutdown()


def test_apply_refuses_to_overwrite_an_existing_out_without_force(tmp_path):
    """I3. `apply()`'s `out` is a real, accepted design (unlike `confirm()`'s
    `out` -- producing an output file IS the operation, and the CLI's own
    `-o` already takes any path). The asymmetry to close is plain data loss
    for the ordinary same-origin user: an existing `out` must be refused
    unless `force` is passed. Byte-level assertion on the victim, not just
    the status code.
    """
    pptx_path = tmp_path / "deck.pptx"
    Presentation().save(str(pptx_path))
    victim = tmp_path / "victim.pptx"
    victim.write_bytes(b"keep me")

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        status, body = _post(port, "/api/apply", {"path": str(pptx_path), "out": str(victim)})
        assert status == 400, body
        assert victim.read_bytes() == b"keep me", (
            "the victim file's bytes changed -- `out` clobbered an existing file without `force`"
        )
    finally:
        srv.shutdown()


def test_apply_with_force_overwrites_an_existing_out(tmp_path):
    """Positive control for I3: `force=True` on a suffix-matched `out` still
    lets the ordinary case (re-running apply on purpose) through.
    """
    pptx_path = tmp_path / "deck2.pptx"
    Presentation().save(str(pptx_path))
    target = tmp_path / "target.pptx"
    target.write_bytes(b"stale")

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        status, body = _post(
            port,
            "/api/apply",
            {"path": str(pptx_path), "out": str(target), "force": True},
        )
        assert status == 200, body
    finally:
        srv.shutdown()

    written = target.read_bytes()
    assert written != b"stale"
    assert written[:2] == b"PK", "expected a real pptx (zip) header"


def test_apply_returns_a_json_safe_result_for_a_deck_with_a_matched_figure(
    tmp_path,
):
    """Regression, found via task-8's own live verification (a real deck
    audited/applied through the actual HTTP route, not just through
    service.apply() in Python -- that direction is what
    test_pptx_powerpoint.py's web-path test covers, and it never touches
    json.dumps() at all).

    `test_apply_with_force_overwrites_an_existing_out` above is this
    route's only other successful-apply test, and its fixture is an EMPTY
    `Presentation()` -- zero pictures, so `rows` is always `[]` and
    `json.dumps()` never sees what `deck.apply()` actually puts in a row:
    a live `Record` instance under `"record"` for every matched picture
    (figcite/deck.py:apply). `service.apply()` used to return that dict
    unchanged, so `/api/apply` 500'd with "Object of type Record is not
    JSON serializable" on any deck that had at least one figure figcite
    could actually cite -- the ordinary case, not an edge case.
    """
    from PIL import Image
    from pptx.util import Inches

    from figcite.provenance import Record, now_stamps

    png = tmp_path / "fig.png"
    Image.new("RGB", (900, 700), (10, 20, 30)).save(png)
    u, loc = now_stamps()
    service.store.register_existing(
        png,
        Record(
            doi="10.9999/regress.1",
            citation="Regression et al. (2026).",
            short_cite="Regression 2026",
            source_kind="pdf-crop",
            captured_utc=u,
            captured_local=loc,
            confirmed=True,
        ),
    )

    pptx_path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(png), Inches(0.5), Inches(0.5), Inches(5))
    prs.save(str(pptx_path))
    out_path = tmp_path / "deck.cited.pptx"

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        status, body = _post(port, "/api/apply", {"path": str(pptx_path), "out": str(out_path)})
        assert status == 200, body
    finally:
        srv.shutdown()

    assert body["out"] == str(out_path)
    assert body["cited"] == 1
    assert out_path.exists()
    # The normalized record must still carry the citation through -- this
    # isn't just "didn't crash", the caller's data has to survive the fix.
    assert body["rows"][0]["record"]["doi"] == "10.9999/regress.1"


def test_apply_refuses_an_out_suffix_that_does_not_match_the_input(tmp_path):
    """I3. A mismatched suffix is refused even when `out` doesn't exist yet
    -- a signal of caller error, not something to "helpfully" honour.
    """
    pptx_path = tmp_path / "deck3.pptx"
    Presentation().save(str(pptx_path))
    bad_out = tmp_path / "sneaky.txt"  # does not exist yet

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        status, body = _post(port, "/api/apply", {"path": str(pptx_path), "out": str(bad_out)})
        assert status == 400, body
    finally:
        srv.shutdown()

    assert not bad_out.exists(), "a file was created at the mismatched-suffix out"


def test_it_rejects_a_forged_host_header_on_get_and_post():
    """I4. DNS-rebinding probe: a page on a public domain that rebinds to
    127.0.0.1 is same-origin to the BROWSER, so its GETs are readable unless
    `Host` is validated too -- `Origin` alone doesn't cover GET (browsers
    don't send `Origin` on a plain navigation/fetch-GET the way they do on
    POST). `/api/pending` leaks capture titles and process names;
    `/api/thumb` leaks actual figure bytes. POSTs must be rejected the same
    way, so both are checked here.
    """
    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        get_resp = _raw_request(
            port,
            b"GET /api/pending HTTP/1.1\r\nHost: evil.example\r\nConnection: close\r\n\r\n",
        )
        assert b" 403 " in _status_of(get_resp), get_resp

        body = b'{"ref": "irrelevant"}'
        post_resp = _raw_request(
            port,
            b"POST /api/skip HTTP/1.1\r\n"
            b"Host: evil.example\r\n"
            b"Content-Type: application/json\r\n"
            + f"Content-Length: {len(body)}\r\n".encode()
            + b"Connection: close\r\n\r\n"
            + body,
        )
        assert b" 403 " in _status_of(post_resp), post_resp
    finally:
        srv.shutdown()


def test_localhost_spelling_is_accepted_on_host_and_origin(tmp_path, monkeypatch):
    """I5. Browsers attach `Origin` to every POST including same-origin
    ones, so a user who navigates to `http://localhost:<port>` (the spelling
    people actually type) got 403 on every mutating request under the
    round-1 Origin check, which only accepted `127.0.0.1`. Both the `Host`
    allowlist (I4) and the `Origin` check must accept `localhost` too.
    """
    _stage(tmp_path, monkeypatch, name="clip-localhost-ok")

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        get_resp = _raw_request(
            port,
            "GET /api/pending HTTP/1.1\r\n".encode()
            + f"Host: localhost:{port}\r\n".encode()
            + b"Connection: close\r\n\r\n",
        )
        assert b" 200 " in _status_of(get_resp), get_resp

        status, body = _post(
            port,
            "/api/skip",
            {"ref": "staged:clip-localhost-ok.png"},
            headers={"Origin": f"http://localhost:{port}"},
        )
        assert status == 200, body
    finally:
        srv.shutdown()


def test_a_non_dict_json_body_gets_400_not_a_crash():
    """M3. `POST /api/skip [1, 2]` used to reach `payload["ref"]` /
    `payload.items()` on a list and 500 with an AttributeError.
    """
    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        status, body = _post(port, "/api/skip", [1, 2])
        assert status == 400, body
        assert "AttributeError" not in json.dumps(body)
    finally:
        srv.shutdown()


def test_make_server_leaves_the_stdlib_class_untouched():
    """M4. `ThreadingHTTPServer.allow_reuse_address = False` (as opposed to
    scoping the override to a private subclass) mutates the actual stdlib
    class object for the rest of the process -- any other
    `ThreadingHTTPServer` anywhere in this process would inherit it, and
    anything that later reset it to True would silently void the
    port-refusal guarantee `test_it_refuses_a_port_already_in_use_rather_
    than_moving` depends on.
    """
    from http.server import ThreadingHTTPServer

    srv = web.make_server(0)
    try:
        assert "allow_reuse_address" not in vars(ThreadingHTTPServer), (
            "make_server() set allow_reuse_address directly on the shared "
            "stdlib ThreadingHTTPServer class instead of a private subclass"
        )
    finally:
        srv.server_close()


def test_concurrent_posts_are_serialized_so_skip_never_double_appends(monkeypatch):
    """M5. The brief's rationale for refusing a taken port was "two servers
    writing one manifest is a corruption path" -- but `ThreadingHTTPServer`
    gives two THREADS in one server the identical hazard, and
    `service.skip`'s check-then-act (`if ref not in _skipped: _skipped.
    append(ref)`) is a real TOCTOU race between them. Force the race
    deterministically by widening the check-then-act window, fire two
    concurrent POSTs for the SAME ref, and assert the module-level lock in
    web.py serializes them so the ref lands in `_skipped` exactly once.
    """

    def _slow_skip(ref):
        if ref not in service._skipped:
            time.sleep(0.2)  # widen the check-then-act window
            service._skipped.append(ref)

    monkeypatch.setattr(service, "skip", _slow_skip)

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        results = []

        def _fire():
            results.append(_post(port, "/api/skip", {"ref": "race-ref"}))

        t1 = threading.Thread(target=_fire)
        t2 = threading.Thread(target=_fire)
        t1.start()
        t2.start()
        t1.join(timeout=5)
        t2.join(timeout=5)

        count = service._skipped.count("race-ref")
        assert count == 1, (
            f"expected exactly one entry, got {count}: {service._skipped!r} "
            "-- two POST threads raced service.skip()"
        )
    finally:
        srv.shutdown()


# --- round 3 security fixes ---------------------------------------------


def test_a_slow_client_does_not_stall_other_clients(tmp_path, monkeypatch):
    """N1. `_LOCK` used to be held across `self.rfile.read(n)` -- unbounded,
    client-paced socket I/O -- and every response write. A lock protects
    SHARED STATE, not the request lifecycle: a client that promises a
    Content-Length and then goes silent parks its handler thread inside
    that read, and (pre-fix) held the lock the whole time, so every OTHER
    mutating request queued behind it. Before the M5 lock existed, a slow
    client only blocked its own thread -- the concurrency fix turned a
    local stall into a global one.

    Reproduces the reviewer's exact probe: a client sends
    `Content-Length: 1000`, delivers 5 bytes, then goes silent. A second,
    complete request must still finish quickly.
    """
    _stage(tmp_path, monkeypatch, name="clip-slow-client-victim")

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    slow_sock = None
    try:
        # First client: claims a 1000-byte body, sends 5 bytes, goes silent.
        # Stays open (never closed, never completes) so the server's
        # `self.rfile.read(n)` truly blocks for the life of this test.
        slow_sock = socket.create_connection(("127.0.0.1", port), timeout=10)
        slow_sock.sendall(
            b"POST /api/skip HTTP/1.1\r\n"
            + f"Host: 127.0.0.1:{port}\r\n".encode()
            + b"Content-Type: application/json\r\n"
            b"Content-Length: 1000\r\n"
            b"\r\n"
            b"12345"  # far short of 1000; the rest never arrives
        )
        # Let the server actually enter the handler and start its read,
        # so the timing assertion below isn't racing the slow client's own
        # connection setup.
        time.sleep(0.3)

        # Second client: a normal, complete request, bounded so a pre-fix
        # stall fails this test within a few seconds instead of hanging the
        # whole run.
        start = time.monotonic()
        status, body = _post(
            port,
            "/api/skip",
            {"ref": "staged:clip-slow-client-victim.png"},
            timeout=8,
        )
        elapsed = time.monotonic() - start

        assert status == 200, body
        assert elapsed < 3, (
            f"second client took {elapsed:.2f}s -- it was blocked behind "
            "the slow client's socket read, meaning the lock (or something "
            "else) is still held across body I/O"
        )
    finally:
        if slow_sock is not None:
            slow_sock.close()
        srv.shutdown()


class _SlowCrossRefResponse:
    """Whatever `fetch_work` needs from a `requests` response, and no more."""

    status_code = 200

    def raise_for_status(self):
        return None

    def json(self):
        return {
            "message": {
                "DOI": "10.1234/slow-service",
                "title": ["A paper CrossRef is slow about"],
                "author": [{"family": "Slow", "given": "S."}],
                "issued": {"date-parts": [[2026]]},
                "container-title": ["Journal of Waiting"],
            }
        }


def test_a_slow_crossref_lookup_does_not_stall_other_clients(tmp_path, monkeypatch):
    """Round 4, I1. The round-3 fix narrowed `_LOCK` off CLIENT-paced I/O but
    left it spanning SERVER-paced I/O: `/api/confirm` held it across
    `service.confirm()`, which reaches `_actions.record_for` ->
    `crossref.record_from_doi` -> `fetch_work` -> `throttled_get` -- a
    25s-timeout request plus a 429 retry that sleeps up to 15s and re-requests,
    so ~65s worst case, against a handler whose own socket `timeout` is 30.

    `test_a_slow_client_does_not_stall_other_clients` above cannot catch this:
    it stalls a CLIENT socket, and the round-3 fix moved the lock off exactly
    that path. This one stalls the SERVICE DEPENDENCY instead -- the network
    call inside `confirm()` -- which is the path still under the lock.

    The two timing assertions are a pair. `elapsed < 2` is the finding; the
    confirm's own `>= STALL` is the control, without which the test would
    pass vacuously if the confirm 400'd before it ever reached CrossRef.
    """
    STALL = 5.0
    _stage(tmp_path, monkeypatch, name="clip-slow-service")

    def _slow_get(url, **kw):
        time.sleep(STALL)
        return _SlowCrossRefResponse()

    monkeypatch.setattr(crossref, "throttled_get", _slow_get)

    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    confirm_result = {}
    try:

        def _confirm():
            began = time.monotonic()
            confirm_result["res"] = _post(
                port,
                "/api/confirm",
                {"ref": "staged:clip-slow-service.png", "doi": "10.1234/slow-service"},
                timeout=30,
            )
            confirm_result["elapsed"] = time.monotonic() - began

        t = threading.Thread(target=_confirm)
        t.start()
        # Let the confirm actually reach the CrossRef call before timing the
        # unrelated request, so this isn't racing the first thread's own setup.
        time.sleep(0.4)

        start = time.monotonic()
        status, body = _post(port, "/api/skip", {"ref": "unrelated-ref"}, timeout=20)
        elapsed = time.monotonic() - start

        assert status == 200, body
        assert elapsed < 2, (
            f"an unrelated POST took {elapsed:.2f}s while a confirm waited on "
            "CrossRef -- the lock is still held across an outbound network call"
        )

        t.join(timeout=30)
    finally:
        srv.shutdown()

    # Control: the confirm really did go through the stalled dependency. If it
    # had failed fast, the assertion above would have proved nothing.
    assert confirm_result.get("elapsed", 0) >= STALL, (
        f"the confirm finished in {confirm_result.get('elapsed')!r}s -- it never "
        "reached the stubbed CrossRef call, so this test proved nothing"
    )
    assert confirm_result["res"][0] == 200, confirm_result["res"]


def test_host_header_comparison_is_case_insensitive():
    """N2. `Host` matching is case-insensitive per RFC 9110; comparing it
    case-sensitively fails closed (not a security hole -- an attacker
    gains nothing from case-shuffling their OWN forged Host) but produces
    an unnecessary 403 for a legitimate client that happens to send a
    differently-cased Host, such as `LOCALHOST`.
    """
    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        resp = _raw_request(
            port,
            b"GET /api/pending HTTP/1.1\r\n"
            + f"Host: LOCALHOST:{port}\r\n".encode()
            + b"Connection: close\r\n\r\n",
        )
        assert b" 200 " in _status_of(resp), resp
    finally:
        srv.shutdown()


# --- round 4, I2: these tests run in the push gate now, which is only safe
# if the fixture that lets them run still blocks everything else -----------


def test_the_network_block_still_refuses_a_non_loopback_address():
    """I2. Every test in this file was `@pytest.mark.live` purely because
    `conftest._block_network` patched `socket.socket.connect`
    unconditionally, and `.githooks/pre-push` runs `pytest -m "not live"` --
    so the whole HTTP security net sat outside the push gate. Exempting
    loopback is what brought them in; this is the check that the exemption
    did not also open the door it exists to keep shut.

    Asserted against the LIVE patched `connect` -- this test is itself
    non-live, so the fixture is active on it and `socket.connect` here IS
    the guard, not a re-implementation of it. No real outbound attempt is
    made: the patch raises before any syscall, which is the point.
    """
    s = socket.socket()
    # A short timeout so that a BROKEN guard (one that lets the address
    # through) fails this test in seconds with a connect error, instead of
    # hanging the run on a real outbound attempt that nothing answers.
    s.settimeout(2)
    try:
        # A public IP literal.
        with pytest.raises(RuntimeError, match="network access blocked"):
            s.connect(("93.184.216.34", 80))
        # A hostname, which `_is_loopback` cannot resolve and so must refuse
        # -- including the loopback-ish spellings, since a name is not an
        # address until something resolves it.
        with pytest.raises(RuntimeError, match="network access blocked"):
            s.connect(("api.crossref.org", 443))
        with pytest.raises(RuntimeError, match="network access blocked"):
            s.connect(("localhost", 80))
    finally:
        s.close()


def test_the_network_block_lets_loopback_through():
    """Positive control for the test above: a negative result needs one, or
    a fixture that blocked EVERYTHING would read as "the guard works" while
    every test in this file failed to connect at all.
    """
    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    listener.listen(1)
    port = listener.getsockname()[1]
    try:
        with socket.create_connection(("127.0.0.1", port), timeout=5) as s:
            assert s.getpeername()[1] == port
    finally:
        listener.close()


def test_own_work_on_a_filed_capture_succeeds_over_a_real_socket(monkeypatch):
    """The behaviour the page test used to encode as "hide the button".

    Asserting the markup lacks a condition proves nothing about whether the
    request works. This drives the actual route, and its negative twin below
    proves the harness can fail.
    """
    from figcite import store
    from figcite.provenance import Record, now_stamps

    u, loc = now_stamps()
    sha = "9a" * 32
    store.put(
        Record(
            sha256=sha,
            dhash="0" * 16,
            source_kind="clipboard",
            confirmed=False,
            captured_utc=u,
            captured_local=loc,
            source_detail={
                "clipboard_capture": {"process": "SnippingTool", "title": "Snipping Tool"}
            },
        )
    )
    srv = web.make_server(0)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        req = urllib.request.Request(
            f"http://127.0.0.1:{port}/api/confirm",
            data=json.dumps({"ref": f"filed:{sha}", "own_work": True}).encode(),
            headers={"Content-Type": "application/json", "Origin": f"http://127.0.0.1:{port}"},
        )
        with urllib.request.urlopen(req) as r:
            assert r.status == 200
            body = json.loads(r.read())
        assert body.get("ok") is True, body

        # Positive control on the OTHER side: the same route must still refuse
        # a request with nothing to confirm by, or the 200 above means nothing.
        bad = urllib.request.Request(
            f"http://127.0.0.1:{port}/api/confirm",
            data=json.dumps({"ref": f"filed:{sha}"}).encode(),
            headers={"Content-Type": "application/json", "Origin": f"http://127.0.0.1:{port}"},
        )
        with pytest.raises(urllib.error.HTTPError) as e:
            urllib.request.urlopen(bad)
        assert e.value.code in (400, 404)
    finally:
        srv.shutdown()
